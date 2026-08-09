"""Execute the locked three-candidate validation without parameter tuning.

The rules live in ``candidate_paths_preregister.json`` and are deliberately
not configurable here.  This module only maps those rules onto the official
AQR and Kenneth French files downloaded for the experiment.  Every selected
input series is checked for duplicate dates, internal gaps, missing values and
unit mistakes before a return is calculated.

Run from the repository root with::

    python -m experiments.candidate_paths \
        --data-dir C:\\path\\to\\wewillsee-candidate-data
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook


LOOKBACK_MONTHS = 36
SLEEVE_VOL_TARGET = 0.10
PORTFOLIO_VOL_TARGET = 0.15
MAX_LEVERAGE = 1.50
ROLLING_MONTHS = 60
COMMON_START = pd.Period("1991-01", freq="M")
DEVELOPMENT_END = pd.Period("2009-12", freq="M")
VALIDATION_START = pd.Period("2010-01", freq="M")
ORIGINAL_END = pd.Period("2009-12", freq="M")

PASS_LOCKED_HISTORICAL = "PASS_LOCKED_HISTORICAL"
FAIL_LOCKED_HISTORICAL = "FAIL_LOCKED_HISTORICAL"
DATA_INCONCLUSIVE = "DATA_INCONCLUSIVE"
DIAGNOSTIC_ONLY = "DIAGNOSTIC_ONLY"
INVALID_PROTOCOL_COVERAGE = "INVALID_PROTOCOL_COVERAGE"

REGION_FILES = {
    "North America": ("north_america_5.zip", "north_america_mom.zip"),
    "Europe": ("europe_5.zip", "europe_mom.zip"),
    "Japan": ("japan_5.zip", "japan_mom.zip"),
    "Asia Pacific ex Japan": (
        "asia_pacific_ex_japan_5.zip",
        "asia_pacific_ex_japan_mom.zip",
    ),
}

TREND_COLUMNS = ("TSMOM", "TSMOM^CM", "TSMOM^EQ", "TSMOM^FI", "TSMOM^FX")
TREND_ASSET_CLASSES = ("TSMOM^CM", "TSMOM^EQ", "TSMOM^FI", "TSMOM^FX")
STYLE_COLUMNS = ("HML", "RMW", "CMA")


class DataValidationError(ValueError):
    """Raised when an input would otherwise be silently omitted or altered."""


@dataclass(frozen=True)
class CostCase:
    trend_annual_drag: float
    equity_style_annual_drag: float
    financing_spread_above_cash: float


@dataclass
class DataBundle:
    trend: pd.DataFrame
    trend_original: pd.DataFrame
    styles: dict[str, pd.DataFrame]
    developed: pd.DataFrame
    common_end: pd.Period
    source_hashes: dict[str, str]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _as_month(value: date | datetime | pd.Timestamp) -> pd.Period:
    return pd.Timestamp(value).to_period("M")


def validate_monthly_frame(
    frame: pd.DataFrame | pd.Series,
    name: str,
    *,
    required_columns: Iterable[str] | None = None,
) -> None:
    """Require a sorted, unique, complete monthly series with finite values."""
    if frame.empty:
        raise DataValidationError(f"{name}: no monthly observations")
    if not isinstance(frame.index, pd.PeriodIndex) or frame.index.freqstr != "M":
        raise DataValidationError(f"{name}: index must be a monthly PeriodIndex")
    if frame.index.has_duplicates:
        duplicates = frame.index[frame.index.duplicated()].astype(str).tolist()
        raise DataValidationError(f"{name}: duplicate months: {duplicates[:5]}")
    if not frame.index.is_monotonic_increasing:
        raise DataValidationError(f"{name}: months are not sorted")
    expected = pd.period_range(frame.index[0], frame.index[-1], freq="M")
    missing_months = expected.difference(frame.index)
    if len(missing_months):
        shown = ", ".join(str(month) for month in missing_months[:10])
        raise DataValidationError(f"{name}: missing month(s): {shown}")

    table = frame.to_frame() if isinstance(frame, pd.Series) else frame
    if required_columns is not None:
        absent = [column for column in required_columns if column not in table.columns]
        if absent:
            raise DataValidationError(f"{name}: missing selected series: {absent}")
        table = table[list(required_columns)]
    numeric = table.apply(pd.to_numeric, errors="coerce")
    bad = numeric.isna() | ~np.isfinite(numeric)
    if bad.to_numpy().any():
        row, column = np.argwhere(bad.to_numpy())[0]
        raise DataValidationError(
            f"{name}: missing/non-finite {numeric.columns[column]} at "
            f"{numeric.index[row]}"
        )


def _trim_warmup(series: pd.Series, name: str) -> pd.Series:
    """Remove only leading warm-up NaNs; reject any later missing result."""
    first = series.first_valid_index()
    if first is None:
        raise DataValidationError(f"{name}: volatility scaling never became available")
    trimmed = series.loc[first:]
    if trimmed.isna().any():
        month = trimmed.index[trimmed.isna()][0]
        raise DataValidationError(f"{name}: internal missing calculated return at {month}")
    validate_monthly_frame(trimmed, name)
    return trimmed.astype(float)


def load_aqr_workbook(path: Path, label: str) -> pd.DataFrame:
    """Parse either the current or original AQR TSMOM workbook by labels."""
    if not path.is_file():
        raise FileNotFoundError(f"{label}: required file not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    matching = [sheet for sheet in workbook.worksheets if "tsmom factor" in sheet.title.lower()]
    if len(matching) != 1:
        raise DataValidationError(
            f"{label}: expected one TSMOM factor sheet, found {[s.title for s in matching]}"
        )
    sheet = matching[0]
    header_row: tuple[Any, ...] | None = None
    rows = list(sheet.iter_rows(values_only=True))
    for row in rows:
        labels = {str(value).strip() for value in row if value is not None}
        if set(TREND_ASSET_CLASSES).issubset(labels):
            header_row = row
            break
    if header_row is None:
        raise DataValidationError(f"{label}: could not locate the four TSMOM asset classes")

    positions = {str(value).strip(): idx for idx, value in enumerate(header_row) if value is not None}
    selected_columns = list(TREND_ASSET_CLASSES)
    if "TSMOM" in positions:
        selected_columns.insert(0, "TSMOM")
    records: list[dict[str, Any]] = []
    for row in rows:
        if not row or not isinstance(row[0], (date, datetime, pd.Timestamp)):
            continue
        record: dict[str, Any] = {"month": _as_month(row[0])}
        for column in selected_columns:
            idx = positions[column]
            value = row[idx] if idx < len(row) else None
            try:
                number = float(value)
            except (TypeError, ValueError) as exc:
                raise DataValidationError(
                    f"{label}: missing/non-numeric {column} at {record['month']}"
                ) from exc
            if not math.isfinite(number):
                raise DataValidationError(f"{label}: non-finite {column} at {record['month']}")
            # AQR workbooks are decimal returns.  A value above 100% is a
            # strong signal that the wrong sheet or units were selected.
            if abs(number) > 1.0:
                raise DataValidationError(
                    f"{label}: implausible decimal return {number} for {column} "
                    f"at {record['month']}"
                )
            record[column] = number
        records.append(record)
    if not records:
        raise DataValidationError(f"{label}: no dated return rows found")
    frame = pd.DataFrame.from_records(records).set_index("month").sort_index()
    frame.index = pd.PeriodIndex(frame.index, freq="M")
    # The lock explicitly permits this one fallback.  It is not used by the
    # official files in this run, which both publish their all-assets factor.
    if "TSMOM" not in frame.columns:
        frame["TSMOM"] = frame[list(TREND_ASSET_CLASSES)].mean(axis=1)
    validate_monthly_frame(frame, label, required_columns=TREND_COLUMNS)
    return frame[list(TREND_COLUMNS)].astype(float)


def load_french_zip(path: Path, required_columns: Iterable[str], label: str) -> pd.DataFrame:
    """Parse only the six-digit monthly section of a French CSV ZIP."""
    if not path.is_file():
        raise FileNotFoundError(f"{label}: required file not found: {path}")
    with zipfile.ZipFile(path) as archive:
        csv_names = [name for name in archive.namelist() if name.lower().endswith(".csv")]
        if len(csv_names) != 1:
            raise DataValidationError(f"{label}: expected one CSV in ZIP, found {csv_names}")
        text = archive.read(csv_names[0]).decode("utf-8-sig")

    parsed = list(csv.reader(io.StringIO(text)))
    required = list(required_columns)
    header: list[str] | None = None
    positions: dict[str, int] = {}
    for row in parsed:
        stripped = [cell.strip() for cell in row]
        if set(required).issubset(set(stripped)):
            header = stripped
            positions = {column: stripped.index(column) for column in required}
            break
    if header is None:
        raise DataValidationError(f"{label}: missing selected series {required}")

    records: list[dict[str, Any]] = []
    for row in parsed:
        if not row or re.fullmatch(r"\s*\d{6}\s*", row[0]) is None:
            continue
        token = row[0].strip()
        month = pd.Period(f"{token[:4]}-{token[4:]}", freq="M")
        record: dict[str, Any] = {"month": month}
        for column, idx in positions.items():
            value = row[idx].strip() if idx < len(row) else ""
            try:
                percentage = float(value)
            except ValueError as exc:
                raise DataValidationError(
                    f"{label}: missing/non-numeric {column} at {month}"
                ) from exc
            if percentage <= -99.0:
                raise DataValidationError(
                    f"{label}: French missing-value sentinel for {column} at {month}"
                )
            if not math.isfinite(percentage):
                raise DataValidationError(f"{label}: non-finite {column} at {month}")
            record[column] = percentage / 100.0
        records.append(record)
    if not records:
        raise DataValidationError(f"{label}: no six-digit monthly rows found")
    frame = pd.DataFrame.from_records(records).set_index("month").sort_index()
    frame.index = pd.PeriodIndex(frame.index, freq="M")
    validate_monthly_frame(frame, label, required_columns=required)
    return frame[required].astype(float)


def load_data(data_dir: Path) -> DataBundle:
    """Load, validate and align every official source required by the lock."""
    expected_files = {
        "aqr_tsmom_current.xlsx",
        "aqr_tsmom_original.xlsx",
        "developed_5.zip",
    }
    for five, momentum in REGION_FILES.values():
        expected_files.update((five, momentum))
    missing_files = sorted(name for name in expected_files if not (data_dir / name).is_file())
    if missing_files:
        raise FileNotFoundError(f"missing required source files: {missing_files}")

    current = load_aqr_workbook(data_dir / "aqr_tsmom_current.xlsx", "AQR current TSMOM")
    original = load_aqr_workbook(data_dir / "aqr_tsmom_original.xlsx", "AQR original TSMOM")
    if original.index[-1] != ORIGINAL_END:
        raise DataValidationError(
            f"AQR original TSMOM: expected frozen endpoint {ORIGINAL_END}, "
            f"found {original.index[-1]}"
        )
    if current.index[0] > ORIGINAL_END + 1:
        raise DataValidationError("AQR current TSMOM does not cover the post-paper extension")

    # The frozen original-paper observations define 1985-2009.  Only dates
    # after the paper endpoint come from AQR's reconstructed current extension.
    trend = pd.concat(
        [original.loc[:ORIGINAL_END], current.loc[ORIGINAL_END + 1 :]],
        axis=0,
    )
    validate_monthly_frame(trend, "spliced AQR TSMOM", required_columns=TREND_COLUMNS)

    styles: dict[str, pd.DataFrame] = {}
    for region, (five_name, mom_name) in REGION_FILES.items():
        five = load_french_zip(
            data_dir / five_name,
            ("Mkt-RF", "HML", "RMW", "CMA", "RF"),
            f"French {region} five factors",
        )
        momentum = load_french_zip(
            data_dir / mom_name,
            ("WML",),
            f"French {region} momentum",
        ).rename(columns={"WML": "Mom"})
        start = max(five.index[0], momentum.index[0])
        end = min(five.index[-1], momentum.index[-1])
        combined = five.loc[start:end].join(momentum.loc[start:end], how="left")
        validate_monthly_frame(
            combined,
            f"French {region} combined styles",
            required_columns=("HML", "RMW", "CMA", "Mom"),
        )
        styles[region] = combined[["HML", "RMW", "CMA", "Mom"]]

    developed = load_french_zip(
        data_dir / "developed_5.zip",
        ("Mkt-RF", "RF"),
        "French Developed benchmark",
    )
    common_end = min(
        [trend.index[-1], developed.index[-1]]
        + [frame.index[-1] for frame in styles.values()]
    )
    if common_end < VALIDATION_START:
        raise DataValidationError(f"common history ends before validation: {common_end}")

    required_index = pd.period_range(COMMON_START, common_end, freq="M")
    for name, frame in [("trend", trend), ("developed", developed), *styles.items()]:
        absent = required_index.difference(frame.index)
        if len(absent):
            shown = ", ".join(str(month) for month in absent[:10])
            raise DataValidationError(f"{name}: missing required common month(s): {shown}")

    source_hashes = {name: _sha256(data_dir / name) for name in sorted(expected_files)}
    return DataBundle(
        trend=trend,
        trend_original=original,
        styles=styles,
        developed=developed,
        common_end=common_end,
        source_hashes=source_hashes,
    )


def causal_vol_scale(
    returns: pd.Series,
    *,
    target: float,
    lookback: int = LOOKBACK_MONTHS,
    cap: float = MAX_LEVERAGE,
) -> pd.Series:
    """Scale from the prior 36 completed months, never the current return."""
    if lookback < 2 or target <= 0 or cap < 0:
        raise ValueError("invalid volatility-scaling parameters")
    validate_monthly_frame(returns, "volatility-scaling input")
    trailing = returns.rolling(lookback, min_periods=lookback).std(ddof=1).shift(1)
    annualized = trailing * math.sqrt(12.0)
    scale = target / annualized
    scale = scale.replace([np.inf, -np.inf], cap).clip(lower=0.0, upper=cap)
    scale.name = "scale"
    return scale


def apply_annual_drag(returns: pd.Series, annual_drag: float) -> pd.Series:
    """Deduct the locked annual drag evenly over twelve monthly marks."""
    if annual_drag < 0:
        raise ValueError("annual drag cannot be negative")
    return returns - annual_drag / 12.0


def scaled_sleeve(
    raw_returns: pd.Series,
    annual_drag: float,
    name: str,
) -> tuple[pd.Series, pd.Series]:
    scale = causal_vol_scale(raw_returns, target=SLEEVE_VOL_TARGET)
    gross = raw_returns * scale
    net = apply_annual_drag(gross, annual_drag)
    return _trim_warmup(net, name), scale


def scaled_sleeve_proportional_drag(
    raw_returns: pd.Series,
    annual_drag: float,
    name: str,
) -> tuple[pd.Series, pd.Series]:
    """Audit-only alternative: multiply the registered drag by actual scale.

    The locked primary subtracts the stated annual drag once from the scaled
    sleeve.  A second reasonable reading is that trading/shorting drag grows
    with the amount of sleeve exposure.  This function records that reading as
    a sensitivity; it does not replace or tune the preregistered primary.
    """
    if annual_drag < 0:
        raise ValueError("annual drag cannot be negative")
    scale = causal_vol_scale(raw_returns, target=SLEEVE_VOL_TARGET)
    gross = raw_returns * scale
    net = gross - scale * (annual_drag / 12.0)
    return _trim_warmup(net, name), scale


def _aligned_frame(series: dict[str, pd.Series], name: str) -> pd.DataFrame:
    starts = []
    ends = []
    for key, value in series.items():
        first = value.first_valid_index()
        last = value.last_valid_index()
        if first is None or last is None:
            raise DataValidationError(f"{name}: {key} has no valid data")
        starts.append(first)
        ends.append(last)
    start, end = max(starts), min(ends)
    frame = pd.concat({key: value.loc[start:end] for key, value in series.items()}, axis=1)
    validate_monthly_frame(frame, name, required_columns=series.keys())
    return frame.astype(float)


def max_drawdown(returns: pd.Series) -> float:
    """Peak-to-trough loss from every monthly mark, including initial cash."""
    if (returns <= -1.0).any():
        month = returns.index[returns <= -1.0][0]
        raise DataValidationError(f"return at or below -100% at {month}")
    wealth = np.concatenate(([1.0], np.cumprod(1.0 + returns.to_numpy(dtype=float))))
    peaks = np.maximum.accumulate(wealth)
    return float(np.min(wealth / peaks - 1.0))


def performance_metrics(returns: pd.Series, risk_free: pd.Series | None = None) -> dict[str, Any]:
    validate_monthly_frame(returns, "performance returns")
    if (returns <= -1.0).any():
        month = returns.index[returns <= -1.0][0]
        raise DataValidationError(f"performance return at or below -100% at {month}")
    months = len(returns)
    total = float(np.prod(1.0 + returns.to_numpy(dtype=float)) - 1.0)
    cagr = float((1.0 + total) ** (12.0 / months) - 1.0)
    volatility = float(returns.std(ddof=1) * math.sqrt(12.0))
    if risk_free is None:
        excess = returns
    else:
        aligned = _aligned_frame({"returns": returns, "risk_free": risk_free}, "Sharpe inputs")
        if not aligned.index.equals(returns.index):
            raise DataValidationError("risk-free series does not cover every performance month")
        excess = aligned["returns"] - aligned["risk_free"]
    excess_std = float(excess.std(ddof=1))
    sharpe = float(excess.mean() / excess_std * math.sqrt(12.0)) if excess_std > 0 else None
    return {
        "start": str(returns.index[0]),
        "end": str(returns.index[-1]),
        "months": months,
        "total_return": total,
        "cagr": cagr,
        "annualized_volatility": volatility,
        "sharpe": sharpe,
        "max_drawdown": max_drawdown(returns),
        "worst_month": float(returns.min()),
    }


def rolling_win_rate(
    returns: pd.Series,
    benchmark: pd.Series | None = None,
    *,
    window: int = ROLLING_MONTHS,
) -> dict[str, Any]:
    values = {"candidate": returns}
    if benchmark is not None:
        values["benchmark"] = benchmark
    frame = _aligned_frame(values, "rolling-window inputs")
    if len(frame) < window:
        raise DataValidationError(
            f"rolling-window inputs: need {window} months, found {len(frame)}"
        )
    candidate_wealth = (1.0 + frame["candidate"]).rolling(window).apply(np.prod, raw=True)
    if benchmark is None:
        valid = candidate_wealth.notna()
        wins = candidate_wealth.loc[valid] > 1.0
    else:
        benchmark_wealth = (1.0 + frame["benchmark"]).rolling(window).apply(np.prod, raw=True)
        valid = candidate_wealth.notna() & benchmark_wealth.notna()
        wins = candidate_wealth.loc[valid] > benchmark_wealth.loc[valid]
    return {"window_months": window, "windows": int(len(wins)), "win_rate": float(wins.mean())}


def _time_slices(series: pd.Series) -> dict[str, pd.Series]:
    slices: dict[str, pd.Series] = {
        "full": series,
        "development_1991_2009": series.loc[:DEVELOPMENT_END],
        "untouched_2010_latest": series.loc[VALIDATION_START:],
    }
    first_decade = (series.index[0].year // 10) * 10
    last_decade = (series.index[-1].year // 10) * 10
    for decade in range(first_decade, last_decade + 1, 10):
        part = series[(series.index.year >= decade) & (series.index.year <= decade + 9)]
        if not part.empty:
            slices[f"decade_{decade}s"] = part
    for required in ("development_1991_2009", "untouched_2010_latest"):
        if slices[required].empty:
            raise DataValidationError(f"required time split is empty: {required}")
    return slices


def metrics_by_period(
    returns: pd.Series,
    risk_free: pd.Series | None = None,
) -> dict[str, dict[str, Any]]:
    output: dict[str, dict[str, Any]] = {}
    for label, part in _time_slices(returns).items():
        rf_part = risk_free.loc[part.index] if risk_free is not None else None
        output[label] = performance_metrics(part, rf_part)
    return output


def _portfolio_and_benchmark(
    p1: pd.Series,
    p2: pd.Series,
    developed: pd.DataFrame,
    financing_spread: float,
    name: str,
) -> dict[str, pd.Series]:
    inputs = _aligned_frame(
        {
            "p1": p1,
            "p2": p2,
            "Mkt-RF": developed["Mkt-RF"],
            "RF": developed["RF"],
        },
        f"{name} portfolio inputs",
    )
    # Factor sleeves are published excess returns.  Add them to the market
    # excess return, scale the risky package, and leave unused capital in RF.
    portfolio_excess_raw = inputs["Mkt-RF"] + 0.50 * inputs["p1"] + 0.50 * inputs["p2"]
    portfolio_scale = causal_vol_scale(portfolio_excess_raw, target=PORTFOLIO_VOL_TARGET)
    benchmark_scale = causal_vol_scale(inputs["Mkt-RF"], target=PORTFOLIO_VOL_TARGET)

    monthly_spread = financing_spread / 12.0
    # At outer scale s, P3 funds s market + 0.5s + 0.5s overlays.  The
    # benchmark funds only s market.  Charge only notional above capital.
    portfolio_funded = (2.0 * portfolio_scale - 1.0).clip(lower=0.0)
    benchmark_funded = (benchmark_scale - 1.0).clip(lower=0.0)
    portfolio = (
        inputs["RF"]
        + portfolio_scale * portfolio_excess_raw
        - portfolio_funded * monthly_spread
    )
    benchmark = (
        inputs["RF"]
        + benchmark_scale * inputs["Mkt-RF"]
        - benchmark_funded * monthly_spread
    )
    portfolio = _trim_warmup(portfolio, f"{name} P3")
    common = _aligned_frame(
        {
            "portfolio": portfolio,
            "benchmark": benchmark,
            "risk_free": inputs["RF"],
            "portfolio_scale": portfolio_scale,
            "benchmark_scale": benchmark_scale,
        },
        f"{name} scaled comparison",
    )
    return {column: common[column] for column in common.columns}


def _portfolio_with_actual_inner_scale_financing(
    p1: pd.Series,
    p2: pd.Series,
    p1_scale: pd.Series,
    p2_scale: pd.Series,
    developed: pd.DataFrame,
    financing_spread: float,
    name: str,
) -> dict[str, pd.Series]:
    """Audit-only funding interpretation that includes inner sleeve scales.

    The locked primary charges financing on ``2 * outer_scale - 1``.  This
    alternative retains the primary sleeve return/cost series but measures
    funded notional as the scaled market plus the two sleeves at their actual
    inner scales.  It is reported only as a sensitivity because the primary
    interpretation was already frozen and run.
    """
    inputs = _aligned_frame(
        {
            "p1": p1,
            "p2": p2,
            "p1_scale": p1_scale,
            "p2_scale": p2_scale,
            "Mkt-RF": developed["Mkt-RF"],
            "RF": developed["RF"],
        },
        f"{name} actual-inner-scale financing inputs",
    )
    portfolio_excess_raw = inputs["Mkt-RF"] + 0.50 * inputs["p1"] + 0.50 * inputs["p2"]
    portfolio_scale = causal_vol_scale(portfolio_excess_raw, target=PORTFOLIO_VOL_TARGET)
    benchmark_scale = causal_vol_scale(inputs["Mkt-RF"], target=PORTFOLIO_VOL_TARGET)
    monthly_spread = financing_spread / 12.0

    risky_notional = portfolio_scale * (
        1.0 + 0.50 * inputs["p1_scale"] + 0.50 * inputs["p2_scale"]
    )
    portfolio_funded = (risky_notional - 1.0).clip(lower=0.0)
    benchmark_funded = (benchmark_scale - 1.0).clip(lower=0.0)
    portfolio = (
        inputs["RF"]
        + portfolio_scale * portfolio_excess_raw
        - portfolio_funded * monthly_spread
    )
    benchmark = (
        inputs["RF"]
        + benchmark_scale * inputs["Mkt-RF"]
        - benchmark_funded * monthly_spread
    )
    portfolio = _trim_warmup(portfolio, f"{name} actual-inner-scale financing P3")
    common = _aligned_frame(
        {
            "portfolio": portfolio,
            "benchmark": benchmark,
            "risk_free": inputs["RF"],
            "portfolio_scale": portfolio_scale,
            "benchmark_scale": benchmark_scale,
            "portfolio_funded_notional": portfolio_funded,
        },
        f"{name} actual-inner-scale financing comparison",
    )
    return {column: common[column] for column in common.columns}


def _metric_advantage(candidate: dict[str, Any], benchmark: dict[str, Any]) -> dict[str, float]:
    candidate_sharpe = candidate["sharpe"]
    benchmark_sharpe = benchmark["sharpe"]
    return {
        "cagr": float(candidate["cagr"] - benchmark["cagr"]),
        "sharpe": float(candidate_sharpe - benchmark_sharpe),
        "max_drawdown": float(candidate["max_drawdown"] - benchmark["max_drawdown"]),
    }


def _case_result(
    p1: pd.Series,
    p2: pd.Series,
    comparison: dict[str, pd.Series],
) -> dict[str, Any]:
    rf = comparison["risk_free"]
    p1_periods = metrics_by_period(p1)
    p2_periods = metrics_by_period(p2)
    p3_periods = metrics_by_period(comparison["portfolio"], rf)
    benchmark_periods = metrics_by_period(comparison["benchmark"], rf)
    advantages = {
        period: _metric_advantage(p3_periods[period], benchmark_periods[period])
        for period in p3_periods
    }
    return {
        "P1": {
            "periods": p1_periods,
            "rolling_60m_positive": rolling_win_rate(p1),
        },
        "P2": {
            "periods": p2_periods,
            "rolling_60m_positive": rolling_win_rate(p2),
        },
        "P3": {
            "periods": p3_periods,
            "benchmark_periods": benchmark_periods,
            "advantages": advantages,
            "rolling_60m_excess": rolling_win_rate(
                comparison["portfolio"], comparison["benchmark"]
            ),
        },
    }


def _standalone_pass(result: dict[str, Any], breadth_positive: int) -> dict[str, Any]:
    full = result["periods"]["full"]
    development = result["periods"]["development_1991_2009"]
    validation = result["periods"]["untouched_2010_latest"]
    checks = {
        "sharpe_at_least_0_50": full["sharpe"] is not None and full["sharpe"] >= 0.50,
        "development_cagr_positive": development["cagr"] > 0.0,
        "validation_cagr_positive": validation["cagr"] > 0.0,
        "rolling_60m_positive_at_least_70pct": (
            result["rolling_60m_positive"]["win_rate"] >= 0.70
        ),
        "max_drawdown_no_worse_than_minus_40pct": full["max_drawdown"] >= -0.40,
        "breadth_at_least_3_of_4": breadth_positive >= 3,
    }
    return {
        "checks": checks,
        "failed_checks": [name for name, passed in checks.items() if not passed],
        "passed": all(checks.values()),
    }


def _p3_pass(base: dict[str, Any], stress: dict[str, Any]) -> dict[str, Any]:
    base_full = base["periods"]["full"]
    base_benchmark = base["benchmark_periods"]["full"]
    base_adv = base["advantages"]["full"]
    stress_adv = stress["advantages"]["full"]
    checks = {
        "cagr_advantage_at_least_5pp": base_adv["cagr"] >= 0.05,
        "sharpe_advantage_at_least_0_25": base_adv["sharpe"] >= 0.25,
        "max_drawdown_no_worse": base_full["max_drawdown"] >= base_benchmark["max_drawdown"],
        "development_excess_cagr_positive": (
            base["advantages"]["development_1991_2009"]["cagr"] > 0.0
        ),
        "validation_excess_cagr_positive": (
            base["advantages"]["untouched_2010_latest"]["cagr"] > 0.0
        ),
        "rolling_60m_excess_at_least_70pct": (
            base["rolling_60m_excess"]["win_rate"] >= 0.70
        ),
        "stress_cagr_advantage_at_least_2pp": stress_adv["cagr"] >= 0.02,
    }
    landslide_checks = {
        "cagr_advantage_at_least_8pp": base_adv["cagr"] >= 0.08,
        "max_drawdown_no_worse": base_full["max_drawdown"] >= base_benchmark["max_drawdown"],
    }
    return {
        "checks": checks,
        "failed_checks": [name for name, passed in checks.items() if not passed],
        "passed": all(checks.values()),
        "landslide_checks": landslide_checks,
        "landslide_failed_checks": [
            name for name, passed in landslide_checks.items() if not passed
        ],
        "landslide": all(landslide_checks.values()),
    }


def _standalone_with_stress(
    base: dict[str, Any],
    stress: dict[str, Any],
    breadth_positive: int,
) -> dict[str, Any]:
    gate = _standalone_pass(base, breadth_positive)
    gate["checks"]["stress_full_cagr_positive"] = (
        stress["periods"]["full"]["cagr"] > 0.0
    )
    gate["failed_checks"] = [
        name for name, passed in gate["checks"].items() if not passed
    ]
    gate["passed"] = all(gate["checks"].values())
    return gate


def protocol_coverage_audit(
    primary_cases: dict[str, Any],
    proportional_cost_cases: dict[str, Any],
) -> dict[str, Any]:
    """Fail closed if any frozen candidate omits the registered common start."""
    required = str(COMMON_START)

    def starts(candidate: str, *, include_benchmark: bool = False) -> dict[str, str]:
        observed: dict[str, str] = {}
        for case_name, cases in primary_cases.items():
            observed[f"locked_primary_{case_name}"] = cases[candidate]["periods"]["full"][
                "start"
            ]
            if include_benchmark:
                observed[f"locked_primary_benchmark_{case_name}"] = cases[candidate][
                    "benchmark_periods"
                ]["full"]["start"]
        for case_name, cases in proportional_cost_cases.items():
            observed[f"proportional_cost_sensitivity_{case_name}"] = cases[candidate][
                "periods"
            ]["full"]["start"]
            if include_benchmark:
                observed[f"proportional_cost_benchmark_{case_name}"] = cases[candidate][
                    "benchmark_periods"
                ]["full"]["start"]
        return observed

    audit: dict[str, Any] = {
        "required_common_start": required,
        "rule": (
            "Every primary series used by a candidate must begin at the frozen "
            "1991-01 common start; a later start is DATA_INCONCLUSIVE, never PASS."
        ),
        "candidates": {},
    }
    for candidate in ("P1", "P2", "P3"):
        observed = starts(candidate, include_benchmark=candidate == "P3")
        exact = all(value == required for value in observed.values())
        reasons = [] if exact else [INVALID_PROTOCOL_COVERAGE]
        audit["candidates"][candidate] = {
            "required_start": required,
            "observed_starts": observed,
            "exact_common_start": exact,
            "valid": exact,
            "reasons": reasons,
        }

    # P3 uses P2 as a required component, so P2 invalidity must propagate even
    # if P3's own output happened to be padded or rebased later.
    p2_valid = audit["candidates"]["P2"]["valid"]
    p3 = audit["candidates"]["P3"]
    p3["inherits_P2_validity"] = p2_valid
    if not p2_valid:
        p3["valid"] = False
        p3["reasons"].append("INHERITS_P2_INVALIDITY")
    return audit


def _honesty_verdict(
    candidate: str,
    evaluated_gate: dict[str, Any],
    coverage: dict[str, Any],
    *,
    proportional_cost_gate: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Separate raw subset diagnostics from the final fail-closed status."""
    output = dict(evaluated_gate)
    output["checks"] = dict(evaluated_gate["checks"])
    output["failed_checks"] = list(evaluated_gate["failed_checks"])
    output["evaluated_subset_gate_passed"] = bool(evaluated_gate["passed"])
    output["protocol_coverage"] = coverage

    alternative_passed = None
    if proportional_cost_gate is not None:
        alternative_passed = bool(proportional_cost_gate["passed"])
        output["cost_interpretation_audit"] = {
            "locked_primary_gate_passed": bool(evaluated_gate["passed"]),
            "drag_proportional_to_actual_scale_gate_passed": alternative_passed,
            "required_for_final_P1_pass": candidate == "P1",
        }

    if not coverage["valid"]:
        output["status"] = DATA_INCONCLUSIVE
        output["status_reason"] = INVALID_PROTOCOL_COVERAGE
        output["status_reasons"] = list(dict.fromkeys(coverage["reasons"]))
        output["passed"] = None
    elif candidate == "P1" and alternative_passed is not True:
        output["status"] = DIAGNOSTIC_ONLY
        output["status_reason"] = "COST_INTERPRETATION_SENSITIVE"
        output["status_reasons"] = ["COST_INTERPRETATION_SENSITIVE"]
        output["passed"] = None
    elif evaluated_gate["passed"]:
        output["status"] = PASS_LOCKED_HISTORICAL
        output["status_reason"] = "ALL_LOCKED_GATES_AND_PROTOCOL_AUDITS_MET"
        output["status_reasons"] = []
        output["passed"] = True
    else:
        output["status"] = FAIL_LOCKED_HISTORICAL
        output["status_reason"] = "LOCKED_GATE_FAILED"
        output["status_reasons"] = ["LOCKED_GATE_FAILED"]
        output["passed"] = False

    if candidate == "P3":
        output["evaluated_subset_landslide"] = bool(evaluated_gate["landslide"])
        if coverage["valid"]:
            output["landslide"] = bool(evaluated_gate["landslide"])
            output["landslide_status"] = (
                "YES" if evaluated_gate["landslide"] else "NO"
            )
        else:
            output["landslide"] = None
            output["landslide_status"] = DATA_INCONCLUSIVE
    return output


def _format_pct(value: float | None) -> str:
    return "n/a" if value is None else f"{100.0 * value:.2f}%"


def _format_num(value: float | None) -> str:
    return "n/a" if value is None else f"{value:.3f}"


def _metric_row(label: str, metrics: dict[str, Any]) -> str:
    return (
        f"| {label} | {metrics['start']} | {metrics['end']} | {metrics['months']} | "
        f"{_format_pct(metrics['cagr'])} | {_format_pct(metrics['annualized_volatility'])} | "
        f"{_format_num(metrics['sharpe'])} | {_format_pct(metrics['max_drawdown'])} | "
        f"{_format_pct(metrics['worst_month'])} |"
    )


def _period_label(label: str) -> str:
    labels = {
        "full": "full evaluable history",
        "development_1991_2009": "evaluable pre-2010 subset (requested 1991-2009)",
        "untouched_2010_latest": "2010+ slice (untouched by this run only)",
    }
    return labels.get(label, label)


def render_report(results: dict[str, Any]) -> str:
    """Render the audited verdict without turning partial coverage green."""
    verdict = results["pass_fail"]
    coverage = results["protocol_audit"]["candidates"]
    cost_sensitivity = results["audit_sensitivities"][
        "cost_drag_proportional_to_actual_scale"
    ]
    financing_sensitivity = results["audit_sensitivities"][
        "financing_notional_includes_inner_sleeve_scale"
    ]
    p2_sensitivity_rate = cost_sensitivity["cases"]["base"]["P2"][
        "rolling_60m_positive"
    ]["win_rate"]
    p3_cost_sensitivity_edge = cost_sensitivity["cases"]["stress"]["P3"][
        "advantages"
    ]["untouched_2010_latest"]["cagr"]
    p3_financing_sensitivity_edge = financing_sensitivity["cases"]["stress"]["P3"][
        "advantages"
    ]["untouched_2010_latest"]["cagr"]

    lines = [
        "# Locked candidate-path validation",
        "",
        f"Generated: {results['metadata']['generated_at_utc']}",
        f"Common published endpoint: {results['metadata']['common_end']}",
        "",
        "This is a one-shot historical test of the committed preregistration. "
        "No result below authorizes live trading or a merge to main.",
        "",
        "## Fail-closed verdict",
        "",
        "| Candidate | Final status | Actual primary start | Reason | Landslide status |",
        "|---|---|---:|---|---|",
        f"| P1 diversified trend | {verdict['P1']['status']} | "
        f"{coverage['P1']['observed_starts']['locked_primary_base']} | "
        f"{verdict['P1']['status_reason']} | n/a |",
        f"| P2 global market-neutral styles | {verdict['P2']['status']} | "
        f"{coverage['P2']['observed_starts']['locked_primary_base']} | "
        f"{verdict['P2']['status_reason']} | n/a |",
        f"| P3 developed core + overlays | {verdict['P3']['status']} | "
        f"{coverage['P3']['observed_starts']['locked_primary_base']} | "
        f"{' and '.join(verdict['P3']['status_reasons'])} | "
        f"{verdict['P3']['landslide_status']} |",
        "",
        "P1 is the only locked historical pass. It starts exactly at 1991-01 and "
        "meets its gate under both the frozen fixed-drag calculation and the audit-only "
        "interpretation that makes drag proportional to actual sleeve scale.",
        "",
        "P2 starts at 1993-11, not the frozen 1991-01 common start. P3 and its benchmark "
        "start at 1996-11, and P3 also depends on the invalid-coverage P2 sleeve. Their raw "
        "numbers remain useful diagnostics, but neither candidate may receive a PASS.",
        "",
        f"On only the evaluable subset, P3's base-cost CAGR advantage was "
        f"{_format_pct(results['cases']['base']['P3']['advantages']['full']['cagr'])} and "
        f"the raw eight-point landslide check was "
        f"{'met' if verdict['P3']['evaluated_subset_landslide'] else 'not met'}. "
        "The locked landslide status is still DATA_INCONCLUSIVE because the frozen start "
        "coverage was not met.",
        "",
        "The 2010+ slice was untouched by this run only. It was not unknown when the ideas "
        "were selected and is not independent of the existing finance literature.",
        "",
        "## Audit-only sensitivities (not the locked primary)",
        "",
        "No parameter, data source, split, cost rate, or trading rule was changed. These are "
        "two explicit readings of how the already-registered cost and funding rates scale.",
        "",
        f"- When sleeve drag is multiplied by actual sleeve scale, P2's base-cost positive "
        f"rate across overlapping 60-month windows is {_format_pct(p2_sensitivity_rate)}, "
        "below its 70% gate.",
        f"- Under that proportional-drag sensitivity, P3's 2010+ stress-cost CAGR edge is "
        f"{_format_pct(p3_cost_sensitivity_edge)} per year.",
        f"- Keeping the locked sleeve-cost calculation but charging funding on the market plus "
        f"the two sleeves at their actual inner scales makes P3's 2010+ stress-cost edge "
        f"{_format_pct(p3_financing_sensitivity_edge)} per year, versus "
        f"{_format_pct(results['cases']['stress']['P3']['advantages']['untouched_2010_latest']['cagr'])} "
        "under the locked primary funding formula.",
        "",
        "These sensitivities are audit warnings, not post-result replacements for the frozen primary.",
        "",
        "## Protocol coverage audit",
        "",
        f"Frozen common start: {results['protocol_audit']['required_common_start']}",
        "",
    ]
    for candidate in ("P1", "P2", "P3"):
        item = coverage[candidate]
        lines.append(
            f"- {candidate}: {'VALID' if item['valid'] else INVALID_PROTOCOL_COVERAGE}; "
            + ", ".join(f"{name}={start}" for name, start in item["observed_starts"].items())
        )

    for case in ("base", "stress"):
        lines.extend(["", f"## Locked primary: {case} costs", ""])
        for candidate in ("P1", "P2"):
            item = results["cases"][case][candidate]
            lines.extend(
                [
                    f"### {candidate}",
                    "",
                    "| Period | Start | End | Months | CAGR | Volatility | Sharpe | Max drawdown | Worst month |",
                    "|---|---:|---:|---:|---:|---:|---:|---:|---:|",
                ]
            )
            for period, metrics in item["periods"].items():
                lines.append(_metric_row(_period_label(period), metrics))
            lines.extend(
                [
                    "",
                    f"Overlapping rolling 60-month positive rate: "
                    f"{_format_pct(item['rolling_60m_positive']['win_rate'])} "
                    f"({item['rolling_60m_positive']['windows']} windows; adjacent windows "
                    "share 59 months and are not independent).",
                    "",
                ]
            )
        p3 = results["cases"][case]["P3"]
        lines.extend(
            [
                "### P3 and identically scaled benchmark",
                "",
                "| Series / period | Start | End | Months | CAGR | Volatility | Sharpe | Max drawdown | Worst month |",
                "|---|---:|---:|---:|---:|---:|---:|---:|---:|",
            ]
        )
        for period, metrics in p3["periods"].items():
            label = _period_label(period)
            lines.append(_metric_row(f"P3 {label}", metrics))
            lines.append(_metric_row(f"Benchmark {label}", p3["benchmark_periods"][period]))
        lines.extend(
            [
                "",
                f"Evaluable-subset full-history CAGR advantage: "
                f"{_format_pct(p3['advantages']['full']['cagr'])}; Sharpe advantage: "
                f"{_format_num(p3['advantages']['full']['sharpe'])}; drawdown advantage: "
                f"{_format_pct(p3['advantages']['full']['max_drawdown'])}.",
                f"Overlapping rolling 60-month benchmark win rate: "
                f"{_format_pct(p3['rolling_60m_excess']['win_rate'])} "
                f"({p3['rolling_60m_excess']['windows']} windows; not independent).",
                "",
            ]
        )

    lines.extend(["## Breadth and independence diagnostics", ""])
    lines.append("### Trend asset classes")
    lines.append("")
    lines.append("| Asset class | Raw annualized average | Base CAGR | Base Sharpe | Base max drawdown | Stress CAGR | Positive |")
    lines.append("|---|---:|---:|---:|---:|---:|---:|")
    for name, item in results["breadth"]["P1_asset_classes"].items():
        lines.append(
            f"| {name} | {_format_pct(item['annualized_average'])} | "
            f"{_format_pct(item['base']['cagr'])} | {_format_num(item['base']['sharpe'])} | "
            f"{_format_pct(item['base']['max_drawdown'])} | "
            f"{_format_pct(item['stress']['cagr'])} | {item['positive']} |"
        )
    lines.extend(["", "### Equity regions after locked primary base costs", ""])
    lines.append("| Region | Start | Base CAGR | Base Sharpe | Base max drawdown | Stress CAGR | Positive |")
    lines.append("|---|---:|---:|---:|---:|---:|---:|")
    for name, item in results["breadth"]["P2_regions"].items():
        lines.append(
            f"| {name} | {item['start']} | {_format_pct(item['base']['cagr'])} | "
            f"{_format_num(item['base']['sharpe'])} | {_format_pct(item['base']['max_drawdown'])} | "
            f"{_format_pct(item['stress']['cagr'])} | {item['positive']} |"
        )
    lines.extend(["", "### Evaluable-subset return correlations", ""])
    corr = results["correlations"]
    labels = list(corr)
    lines.append("| | " + " | ".join(labels) + " |")
    lines.append("|---|" + "---:|" * len(labels))
    for row in labels:
        lines.append("| " + row + " | " + " | ".join(f"{corr[row][col]:.3f}" for col in labels) + " |")

    lines.extend(["", "## Locked gate diagnostics", ""])
    for candidate in ("P1", "P2", "P3"):
        item = verdict[candidate]
        lines.extend(
            [
                f"### {candidate}",
                "",
                f"Final status: {item['status']} ({item['status_reason']}).",
                f"Raw gate on the evaluable subset: "
                f"{'met' if item['evaluated_subset_gate_passed'] else 'not met'}.",
                "",
            ]
        )
        for check, passed in item["checks"].items():
            lines.append(f"- {'MET' if passed else 'NOT MET'} - {check}")
        if "cost_interpretation_audit" in item:
            alternative = item["cost_interpretation_audit"]
            lines.extend(
                [
                    "",
                    "Cost-interpretation audit: locked primary gate "
                    f"{'met' if alternative['locked_primary_gate_passed'] else 'not met'}; "
                    "proportional-drag sensitivity gate "
                    f"{'met' if alternative['drag_proportional_to_actual_scale_gate_passed'] else 'not met'}.",
                ]
            )
        if candidate == "P3":
            lines.extend(
                [
                    "",
                    f"Locked landslide status: {item['landslide_status']}.",
                    "Evaluated-subset landslide check: "
                    f"{'YES' if item['evaluated_subset_landslide'] else 'NO'}.",
                ]
            )
        lines.append("")

    older = results["older_trend_stress"]
    lines.extend(
        [
            "## Frozen original-paper older trend diagnostic",
            "",
            "| Series | Start | End | Months | CAGR | Volatility | Sharpe | Max drawdown | Worst month |",
            "|---|---:|---:|---:|---:|---:|---:|---:|---:|",
            _metric_row(
                "P1 base costs (1985-1990 request; effective after warm-up)",
                older["base"],
            ),
            _metric_row(
                "P1 stress costs (1985-1990 request; effective after warm-up)",
                older["stress"],
            ),
            "",
            "The 36-month causal volatility lookback leaves only 1988-01 through "
            "1990-12 evaluable in the requested 1985-1990 slice.",
            "",
            "## Calculation and funding assumptions",
            "",
            "- AQR original-paper returns are used through 2009-12; the current workbook is used only from 2010-01 onward.",
            "- French percentages are converted to decimal returns exactly once.",
            "- Locked primary sleeve costs divide the annual drag by 12 and subtract it once after sleeve scaling.",
            "- The audit-only proportional-drag sensitivity multiplies that same rate by actual sleeve scale; rates are unchanged.",
            "- Volatility at month t uses only t-36 through t-1; leverage is capped at 1.50.",
            "- P3 and its benchmark keep unallocated capital at the published risk-free return.",
            "- Locked primary P3 financing is max(2 x outer scale - 1, 0); it does not include inner sleeve scales.",
            "- The audit-only financing sensitivity uses max(outer scale x (1 + 0.5 x P1 scale + 0.5 x P2 scale) - 1, 0).",
            "- Benchmark financing is max(benchmark scale - 1, 0) in both readings.",
            "- Monthly wealth, including the initial cash mark, is used for drawdown.",
            "- Rolling 60-month windows overlap by 59 months, so their win-rate observations are not independent.",
            "",
            "## Limits",
            "",
            "These are hypothetical research factors, not executable broker fills. Public series "
            "can be revised, and short borrow, turnover, market impact, and capacity are simplified. "
            "The 2010+ slice was untouched by this run only, not unknown when the ideas were chosen. "
            "P1's historical pass still requires the frozen 6-12 month paper-forward gate; P2 and P3 "
            "first require a valid rerun with complete frozen-start coverage.",
            "",
        ]
    )
    return "\n".join(lines)


def execute(spec_path: Path, data_dir: Path, output_dir: Path) -> dict[str, Any]:
    with spec_path.open("r", encoding="utf-8") as handle:
        spec = json.load(handle)
    if spec.get("status") != "LOCKED_BEFORE_DATA_DOWNLOAD":
        raise DataValidationError("preregistration is not marked LOCKED_BEFORE_DATA_DOWNLOAD")

    shared = spec["shared_rules"]
    expected_shared = {
        "frequency": "monthly",
        "volatility_estimate": "36 trailing completed months, shifted one month so the current return is never used",
        "sleeve_volatility_target": SLEEVE_VOL_TARGET,
        "whole_portfolio_volatility_target": PORTFOLIO_VOL_TARGET,
        "maximum_scaling_leverage": MAX_LEVERAGE,
        "minimum_scaling_leverage": 0.0,
        "missing_month_rule": "Do not silently drop a selected series. Stop the run and report the missing month or series.",
        "compounding": "geometric from a monthly wealth curve",
        "drawdown": "measured from every monthly portfolio mark, never only at trade exits",
        "rolling_consistency_window_months": ROLLING_MONTHS,
    }
    for key, expected in expected_shared.items():
        if shared.get(key) != expected:
            raise DataValidationError(
                f"locked parameter {key!r} changed: expected {expected!r}, got {shared.get(key)!r}"
            )

    expected_data = {
        "common_start": "1991-01",
        "development_end": "2009-12",
        "untouched_validation_start": "2010-01",
        "older_trend_stress_period": "1985-01 through 1990-12",
    }
    for key, expected in expected_data.items():
        if spec["data"].get(key) != expected:
            raise DataValidationError(
                f"locked data boundary {key!r} changed: expected {expected!r}, "
                f"got {spec['data'].get(key)!r}"
            )
    if [candidate.get("id") for candidate in spec.get("candidates", [])] != ["P1", "P2", "P3"]:
        raise DataValidationError("locked candidate set is not exactly P1, P2, P3")

    expected_costs = {
        "base_costs": {
            "trend_annual_drag": 0.01,
            "equity_style_annual_drag": 0.02,
            "financing_spread_above_cash": 0.015,
        },
        "stress_costs": {
            "trend_annual_drag": 0.02,
            "equity_style_annual_drag": 0.03,
            "financing_spread_above_cash": 0.03,
        },
    }
    for key, expected in expected_costs.items():
        if shared.get(key) != expected:
            raise DataValidationError(
                f"locked cost case {key!r} changed: expected {expected!r}, got {shared.get(key)!r}"
            )

    cases = {
        "base": CostCase(**shared["base_costs"]),
        "stress": CostCase(**shared["stress_costs"]),
    }
    data = load_data(data_dir)
    trend_raw = data.trend["TSMOM"].loc[: data.common_end]

    region_raw: dict[str, pd.Series] = {}
    for region, frame in data.styles.items():
        region_raw[region] = frame[["HML", "RMW", "CMA", "Mom"]].mean(axis=1)
        validate_monthly_frame(region_raw[region], f"{region} style sleeve")
    global_style_raw = _aligned_frame(region_raw, "global regional sleeves").mean(axis=1)

    case_results: dict[str, Any] = {}
    monthly_columns: dict[str, pd.Series] = {}
    case_series: dict[str, dict[str, Any]] = {}
    for case_name, costs in cases.items():
        p1, p1_scale = scaled_sleeve(
            trend_raw,
            costs.trend_annual_drag,
            f"P1 {case_name}",
        )
        p2, p2_scale = scaled_sleeve(
            global_style_raw,
            costs.equity_style_annual_drag,
            f"P2 {case_name}",
        )
        p1 = p1.loc[COMMON_START : data.common_end]
        p2 = p2.loc[COMMON_START : data.common_end]
        comparison = _portfolio_and_benchmark(
            p1,
            p2,
            data.developed.loc[: data.common_end],
            costs.financing_spread_above_cash,
            case_name,
        )
        case_results[case_name] = _case_result(p1, p2, comparison)
        case_series[case_name] = {
            "p1": p1,
            "p2": p2,
            "p1_scale": p1_scale.loc[p1.index],
            "p2_scale": p2_scale.loc[p2.index],
            "comparison": comparison,
        }
        monthly_columns.update(
            {
                f"P1_{case_name}": p1,
                f"P1_scale_{case_name}": p1_scale.loc[p1.index],
                f"P2_{case_name}": p2,
                f"P2_scale_{case_name}": p2_scale.loc[p2.index],
                f"P3_{case_name}": comparison["portfolio"],
                f"Benchmark_{case_name}": comparison["benchmark"],
                f"P3_scale_{case_name}": comparison["portfolio_scale"],
                f"Benchmark_scale_{case_name}": comparison["benchmark_scale"],
            }
        )
        if case_name == "base":
            monthly_columns["Risk_free"] = comparison["risk_free"]

    # Audit-only cost reading: the exact registered drag is multiplied by the
    # sleeve's actual causal scale.  The locked primary above remains intact.
    proportional_cost_case_results: dict[str, Any] = {}
    for case_name, costs in cases.items():
        p1, p1_scale = scaled_sleeve_proportional_drag(
            trend_raw,
            costs.trend_annual_drag,
            f"P1 {case_name} proportional-drag sensitivity",
        )
        p2, p2_scale = scaled_sleeve_proportional_drag(
            global_style_raw,
            costs.equity_style_annual_drag,
            f"P2 {case_name} proportional-drag sensitivity",
        )
        p1 = p1.loc[COMMON_START : data.common_end]
        p2 = p2.loc[COMMON_START : data.common_end]
        comparison = _portfolio_and_benchmark(
            p1,
            p2,
            data.developed.loc[: data.common_end],
            costs.financing_spread_above_cash,
            f"{case_name} proportional-drag sensitivity",
        )
        proportional_cost_case_results[case_name] = _case_result(p1, p2, comparison)
        monthly_columns.update(
            {
                f"P1_{case_name}_proportional_drag_sensitivity": p1,
                f"P2_{case_name}_proportional_drag_sensitivity": p2,
                f"P3_{case_name}_proportional_drag_sensitivity": comparison["portfolio"],
            }
        )

    # Separate audit-only funding reading.  Keep the locked sleeve costs but
    # include their inner scales in funded notional above one.
    financing_sensitivity_case_results: dict[str, Any] = {}
    for case_name, costs in cases.items():
        primary = case_series[case_name]
        comparison = _portfolio_with_actual_inner_scale_financing(
            primary["p1"],
            primary["p2"],
            primary["p1_scale"],
            primary["p2_scale"],
            data.developed.loc[: data.common_end],
            costs.financing_spread_above_cash,
            f"{case_name} financing sensitivity",
        )
        financing_sensitivity_case_results[case_name] = _case_result(
            primary["p1"], primary["p2"], comparison
        )
        monthly_columns[
            f"P3_{case_name}_actual_inner_scale_financing_sensitivity"
        ] = comparison["portfolio"]

    p1_breadth: dict[str, Any] = {}
    for column in TREND_ASSET_CLASSES:
        selected = data.trend[column].loc[COMMON_START : data.common_end]
        annual_average = float(selected.mean() * 12.0)
        base_net, _ = scaled_sleeve(
            data.trend[column].loc[: data.common_end],
            cases["base"].trend_annual_drag,
            f"{column} base breadth sleeve",
        )
        stress_net, _ = scaled_sleeve(
            data.trend[column].loc[: data.common_end],
            cases["stress"].trend_annual_drag,
            f"{column} stress breadth sleeve",
        )
        p1_breadth[column] = {
            "annualized_average": annual_average,
            "positive": bool(annual_average > 0.0),
            "base": performance_metrics(base_net.loc[COMMON_START : data.common_end]),
            "stress": performance_metrics(stress_net.loc[COMMON_START : data.common_end]),
        }
    p1_positive = sum(item["positive"] for item in p1_breadth.values())

    p2_breadth: dict[str, Any] = {}
    for region, raw in region_raw.items():
        net, _ = scaled_sleeve(
            raw,
            cases["base"].equity_style_annual_drag,
            f"{region} base breadth sleeve",
        )
        stress_net, _ = scaled_sleeve(
            raw,
            cases["stress"].equity_style_annual_drag,
            f"{region} stress breadth sleeve",
        )
        net = net.loc[COMMON_START : data.common_end]
        stress_net = stress_net.loc[COMMON_START : data.common_end]
        metrics = performance_metrics(net)
        p2_breadth[region] = {
            "start": metrics["start"],
            "end": metrics["end"],
            "base": metrics,
            "stress": performance_metrics(stress_net),
            "positive": bool(metrics["cagr"] > 0.0),
        }
    p2_positive = sum(item["positive"] for item in p2_breadth.values())

    base_series = case_series["base"]
    base_comparison = base_series["comparison"]
    correlation_input = _aligned_frame(
        {
            "P1": base_series["p1"],
            "P2": base_series["p2"],
            "Developed benchmark excess": (
                base_comparison["benchmark"] - base_comparison["risk_free"]
            ),
        },
        "correlation inputs",
    )
    correlations = correlation_input.corr().to_dict()

    # The older test must stay on the frozen workbook, never AQR's later
    # reconstructed history.  Warm-up is retained before slicing to 1990-12.
    older_p1, _ = scaled_sleeve(
        data.trend_original["TSMOM"],
        cases["base"].trend_annual_drag,
        "P1 original-paper older stress",
    )
    older_p1_stress, _ = scaled_sleeve(
        data.trend_original["TSMOM"],
        cases["stress"].trend_annual_drag,
        "P1 original-paper older stress-cost case",
    )
    older_p1 = older_p1.loc[: pd.Period("1990-12", freq="M")]
    older_p1_stress = older_p1_stress.loc[: pd.Period("1990-12", freq="M")]
    older_metrics = {
        "base": performance_metrics(older_p1),
        "stress": performance_metrics(older_p1_stress),
    }

    primary_p1_gate = _standalone_with_stress(
        case_results["base"]["P1"], case_results["stress"]["P1"], p1_positive
    )
    primary_p2_gate = _standalone_with_stress(
        case_results["base"]["P2"], case_results["stress"]["P2"], p2_positive
    )
    primary_p3_gate = _p3_pass(
        case_results["base"]["P3"], case_results["stress"]["P3"]
    )
    proportional_p1_gate = _standalone_with_stress(
        proportional_cost_case_results["base"]["P1"],
        proportional_cost_case_results["stress"]["P1"],
        p1_positive,
    )
    proportional_p2_gate = _standalone_with_stress(
        proportional_cost_case_results["base"]["P2"],
        proportional_cost_case_results["stress"]["P2"],
        p2_positive,
    )
    proportional_p3_gate = _p3_pass(
        proportional_cost_case_results["base"]["P3"],
        proportional_cost_case_results["stress"]["P3"],
    )

    protocol_audit = protocol_coverage_audit(
        case_results, proportional_cost_case_results
    )
    p1_pass = _honesty_verdict(
        "P1",
        primary_p1_gate,
        protocol_audit["candidates"]["P1"],
        proportional_cost_gate=proportional_p1_gate,
    )
    p2_pass = _honesty_verdict(
        "P2",
        primary_p2_gate,
        protocol_audit["candidates"]["P2"],
        proportional_cost_gate=proportional_p2_gate,
    )
    p3_pass = _honesty_verdict(
        "P3",
        primary_p3_gate,
        protocol_audit["candidates"]["P3"],
        proportional_cost_gate=proportional_p3_gate,
    )

    results: dict[str, Any] = {
        "metadata": {
            "generated_at_utc": datetime.now(timezone.utc).isoformat(),
            "preregistered_at": spec["registered_at"],
            "preregistration_status": spec["status"],
            "preregistration_sha256": _sha256(spec_path),
            "common_start": str(COMMON_START),
            "common_end": str(data.common_end),
            "trend_splice": "AQR original through 2009-12; AQR current extension from 2010-01",
            "annual_cost_conversion": (
                "locked primary divides annual drag by 12 and deducts it once after "
                "sleeve scaling"
            ),
            "slice_language": (
                "2010+ was untouched by this run only; it was not unknown when the "
                "ideas were selected or in the finance literature"
            ),
            "rolling_window_dependence": (
                "60-month windows overlap by 59 months and are not independent"
            ),
            "source_sha256": data.source_hashes,
        },
        "cases": case_results,
        "protocol_audit": protocol_audit,
        "audit_sensitivities": {
            "cost_drag_proportional_to_actual_scale": {
                "status": "AUDIT_ONLY_NOT_LOCKED_PRIMARY",
                "assumption": (
                    "multiply each unchanged registered sleeve drag by the sleeve's "
                    "actual causal scaling factor"
                ),
                "cases": proportional_cost_case_results,
                "gate_diagnostics": {
                    "P1": proportional_p1_gate,
                    "P2": proportional_p2_gate,
                    "P3": proportional_p3_gate,
                },
            },
            "financing_notional_includes_inner_sleeve_scale": {
                "status": "AUDIT_ONLY_NOT_LOCKED_PRIMARY",
                "assumption": (
                    "retain locked sleeve returns and costs; funded P3 notional is "
                    "outer_scale * (1 + 0.5 * P1_scale + 0.5 * P2_scale) - 1"
                ),
                "cases": financing_sensitivity_case_results,
            },
        },
        "breadth": {
            "P1_asset_classes": p1_breadth,
            "P1_positive_count": int(p1_positive),
            "P2_regions": p2_breadth,
            "P2_positive_count": int(p2_positive),
        },
        "correlations": correlations,
        "older_trend_stress": older_metrics,
        "pass_fail": {"P1": p1_pass, "P2": p2_pass, "P3": p3_pass},
        "promotion_gate": spec["promotion_gate"],
    }

    output_dir.mkdir(parents=True, exist_ok=True)
    result_path = output_dir / "candidate_paths_results.json"
    report_path = output_dir / "candidate_paths_report.md"
    series_path = output_dir / "candidate_paths_monthly.csv"
    result_path.write_text(json.dumps(results, indent=2, sort_keys=True), encoding="utf-8")
    report_path.write_text(render_report(results), encoding="utf-8")
    monthly = pd.concat(monthly_columns, axis=1).sort_index()
    monthly.index = monthly.index.astype(str)
    monthly.index.name = "month"
    monthly.to_csv(series_path, float_format="%.12g")
    return results


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Run the immutable three-path candidate validation",
    )
    parser.add_argument(
        "--spec",
        type=Path,
        default=Path(__file__).with_name("candidate_paths_preregister.json"),
    )
    parser.add_argument("--data-dir", required=True, type=Path)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(__file__).with_name("results"),
    )
    args = parser.parse_args()
    results = execute(args.spec, args.data_dir, args.output_dir)
    verdict = results["pass_fail"]
    print(
        "Locked result: "
        f"P1={verdict['P1']['status']}, "
        f"P2={verdict['P2']['status']}, "
        f"P3={verdict['P3']['status']}, "
        f"landslide={verdict['P3']['landslide_status']}"
    )


if __name__ == "__main__":
    main()
