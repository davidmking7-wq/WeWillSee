"""picks.xlsx — the running scorecard, in plain language.

Sheets:
- Picks: one row per pick. New picks appended each run; old picks re-priced
  and resolved (HIT once +5% is reached before the deadline, MISS after).
- Track Record: rebuilt each update — how the tool is actually doing vs what
  it predicted.
- How To Read This: every column explained simply.

Every row carries the signal-engine version that produced it ("Engine").
Workbooks created before the column existed are migrated on load: the column
is appended and old rows are stamped "v1", so hit-rates never silently mix
engines. The Track Record breaks results out by engine.
"""
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config

HEADERS = ["Date Picked", "Stock", "Company", "Industry", "List",
           "Price Then", "Min Goal (+5%)", "Deadline",
           # what history says about stocks that looked like this
           "Chance +5%", "Chance +10%", "Usual Gain %", "Usual Peak %",
           "Analysts' Guess % (corrected)",
           "Usual Days to +5%", "Chance of -5% Dip First",
           "Avg Stock Chance +5%", "Overall Score", "Gain Score", "Confidence",
           "Why (short)",
           # what actually happened
           "Result", "Price Now", "Gain So Far %", "Best So Far %",
           "Hit Date", "Last Checked",
           # which version of the scanner made this pick
           "Engine",
           # columns added later — kept at the end so old workbooks migrate
           # by appending, never by shifting existing data
           "Chance +15%", "Earnings Before Deadline", "Sell Signal",
           "Sell Below (Disaster)"]
COL = {h: i + 1 for i, h in enumerate(HEADERS)}

# old header -> new header (renamed in place on migration; renames never
# shift columns, so they're always safe)
RENAMES = {"Goal (+5%)": "Min Goal (+5%)"}

FILL_HEAD = PatternFill("solid", fgColor="1F3B57")
FILL_PRED = PatternFill("solid", fgColor="2E4E6E")
FILL_HIT = PatternFill("solid", fgColor="C6EFCE")
FILL_MISS = PatternFill("solid", fgColor="FFC7CE")

PRED_COLS = ("Chance +5%", "Chance +10%", "Chance +15%", "Usual Gain %",
             "Usual Peak %", "Analysts' Guess % (corrected)",
             "Usual Days to +5%", "Chance of -5% Dip First",
             "Avg Stock Chance +5%", "Overall Score", "Gain Score")

ABOUT = """How to read this scorecard

Every time the scout runs, it adds its picks here and re-checks the old ones.
A pick is a HIT if the stock's closing price gained 5% or more at any point
before the Deadline (about 2 months). After the Deadline, it's a MISS.

IMPORTANT: 5% is the MINIMUM bar, not the goal. The tool ranks picks by how
far and how fast they usually run (see Usual Peak %, Chance +10%, Chance
+15%), and a pick that HITs keeps being re-priced until its Deadline — so
"Best So Far %" shows how far the winner actually ran, not just that it
cleared the bar.

There are two lists:
- "Best Overall": the best balance of chance, size of gain, speed and safety.
- "Biggest Gain": the biggest expected gains, but only stocks with at least
  a 62% chance of reaching +5%. Some runs this list is short or empty —
  that means nothing honest qualified, not that the tool forgot.
- "Both" means the stock made both lists.

What the prediction columns mean (all measured from 10 years of stocks that
looked the same — same pattern, same energy level, same market mood):

Chance +5%        How often such stocks gained at least 5% within 2 months.
Chance +10%       How often they gained at least 10%.
Chance +15%       How often they gained at least 15%.
Usual Gain %      Where they typically ENDED after 2 months (often below 5 —
                  that's the honest truth, not a mistake).
Usual Peak %      The best gain they typically reached at some point (the
                  ranking uses the AVERAGE peak, which gives extra credit to
                  groups with big winners — because 5% is only the minimum).
Analysts' Guess % (corrected)  What Wall Street analysts' price targets say
                  the stock will do over ONE YEAR — after correcting their
                  well-documented over-optimism. Analysts overshoot by about
                  7-12 points on average, so we shrink their claim:
                  corrected = 5% + 0.22 x their claimed gain. "n/r" means
                  their targets disagree so widely (or are so stale) that the
                  number isn't reliable — research shows widely-scattered
                  targets actually predict BADLY. Note this is a 1-YEAR view
                  shown for context; the pick itself is a 2-month idea chosen
                  from history, and a giant claimed upside is a warning sign,
                  not a promise.
Usual Days to +5% How many trading days it usually took to reach +5%.
Chance of -5% Dip First  How often they fell 5% before rising 5%. Important!
Avg Stock Chance +5%     How often ANY normal S&P stock reached +5% in the
                  same kind of market. In good times most do — a pick is only
                  special if its chance beats this number.
Overall Score     Chance x size of gain x speed x safety, in one number.
                  Higher = better. Only compare within the same run.
Gain Score        Usual peak x chance of +10% — the "how big" ranking.
Confidence        A = proven edge over the average stock. B = modest edge.
                  C = no proven edge (kept only if something else is special).
Engine            Which version of the scanner chose the pick (v1, v3...).
                  Versions are scored separately in the Track Record so an
                  upgrade can't hide behind the old version's results.
Earnings Before Deadline  The company's next quarterly report date, if it
                  lands before the Deadline. Know the date — a report can
                  wreck a single pick overnight — but on real SEC data,
                  picks WITH a report ahead in their window were the
                  BETTER group (the report is the fuel for the +5% move).
                  What the tool avoids instead is stocks that JUST
                  reported: with the report behind them they reached +5%
                  only ~40-47% of the time vs ~62% normally, so they get
                  pushed down the lists automatically.
Sell Signal       When to sell, updated every run — and specific to EACH
                  stock, not one wide rule. Ten years of testing say
                  ordinary stop-losses make results WORSE (stocks dip and
                  recover too often, and crashes gap right through stops).
                  The rules that survived testing:
                  1) Disaster stop, sized to the stock: if it closes below
                     "Sell Below (Disaster)" — 2x that stock's own normal
                     2-month move under the buy price — sell; the pattern
                     is broken. A calm stock gets a tighter level, a lively
                     one gets more room. Rarely fires; caps catastrophes.
                  2) Otherwise no stop — sell at the Deadline.
                  3) Once it has touched +5%, never let it become a loss:
                     sell on any close back at/below your buy price.
                     (A tighter "8% below its peak" version was tested and
                     sold winners that kept running — rejected.)
                  This column shows the live instruction with exact price
                  levels. Protection buys smaller losses, not bigger gains
                  (roughly free if you re-invest the freed cash).
Sell Below (Disaster)  This stock's own disaster price, computed on the day
                  it was picked from how much it normally moves.

What-happened columns:
Gain So Far %     Where the stock is now vs the pick price.
Best So Far %     The best it has reached since picked — shows near-misses.

Honesty notes: numbers come from 2016-2026 — mostly good years for stocks;
the tool refuses to give numbers when history is too thin; reaching +5% at
some point is NOT the same as ending with a profit. This is a research
scorecard, NOT financial advice. Nothing is ever bought automatically.
"""


def _migrate(wb: Workbook) -> None:
    """Rename retitled headers in place, then append any headers added since
    the workbook was created (new columns only ever go at the END of HEADERS
    so old data never shifts). Rows written before the Engine column existed
    are stamped 'v1'."""
    ws = wb["Picks"]
    for i in range(1, ws.max_column + 1):
        old = ws.cell(row=1, column=i).value
        if old in RENAMES:
            ws.cell(row=1, column=i, value=RENAMES[old])
    existing = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    for h in HEADERS:
        if h in existing:
            continue
        cidx = COL[h]
        cell = ws.cell(row=1, column=cidx, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = FILL_PRED if h in PRED_COLS else FILL_HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(cidx)].width = 11
        if h == "Engine":
            for row in range(2, ws.max_row + 1):
                if ws.cell(row=row, column=COL["Stock"]).value:
                    ws.cell(row=row, column=cidx, value="v1")


def _load() -> Workbook:
    wb = load_workbook(config.EXCEL_PATH)
    _migrate(wb)
    return wb


def _ensure() -> Workbook:
    if config.EXCEL_PATH.exists():
        return _load()
    wb = Workbook()
    ws = wb.active
    ws.title = "Picks"
    ws.append(HEADERS)
    for i, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = FILL_PRED if h in PRED_COLS else FILL_HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    widths = {"Date Picked": 11, "Stock": 7, "Company": 22, "Industry": 18,
              "List": 12, "Why (short)": 60, "Result": 8, "Deadline": 11,
              "Hit Date": 11, "Last Checked": 12, "Engine": 8}
    for h in HEADERS:
        ws.column_dimensions[get_column_letter(COL[h])].width = widths.get(h, 11)
    ws.freeze_panes = "C2"
    wb.create_sheet("Track Record")
    about = wb.create_sheet("How To Read This")
    for i, line in enumerate(ABOUT.strip().splitlines(), 1):
        about.cell(row=i, column=1, value=line)
    about.column_dimensions["A"].width = 100
    return wb


def append_picks(rows: list[dict]) -> int:
    """rows: dicts keyed like HEADERS (missing keys become blank)."""
    wb = _ensure()
    ws = wb["Picks"]
    for r in rows:
        ws.append([r.get(h, "") for h in HEADERS])
        ws.cell(row=ws.max_row, column=COL["Why (short)"]).alignment = \
            Alignment(wrap_text=True)
    _rebuild_track_record(wb)
    wb.save(config.EXCEL_PATH)
    return len(rows)


def open_picks() -> list[dict]:
    """Rows that still need re-pricing, as dicts including their sheet row
    number: every OPEN pick, plus HIT picks whose Deadline hasn't passed —
    +5% is the minimum bar, not the goal, so winners keep being tracked to
    the deadline to show how far they actually ran."""
    if not config.EXCEL_PATH.exists():
        return []
    wb = _load()
    ws = wb["Picks"]
    today = date.today()
    out = []
    for row in range(2, ws.max_row + 1):
        result = ws.cell(row=row, column=COL["Result"]).value
        if result == "HIT":
            dl = ws.cell(row=row, column=COL["Deadline"]).value
            try:
                dl = dl.date() if hasattr(dl, "date") else date.fromisoformat(str(dl)[:10])
            except ValueError:
                dl = None
            if dl is None or today > dl:
                continue
        elif result != "OPEN":
            continue
        d = {h: ws.cell(row=row, column=COL[h]).value for h in HEADERS}
        d["_row"] = row
        out.append(d)
    return out


def resolve(updates: list[dict]) -> None:
    """updates: {_row, Price Now, Gain So Far %, Best So Far %, Result, Hit Date?}."""
    wb = _load()
    ws = wb["Picks"]
    today = str(date.today())
    for u in updates:
        row = u["_row"]
        for h in ("Price Now", "Gain So Far %", "Best So Far %", "Result",
                  "Hit Date", "Sell Signal"):
            if h in u and u[h] is not None:
                ws.cell(row=row, column=COL[h], value=u[h])
        ws.cell(row=row, column=COL["Last Checked"], value=today)
        result = u.get("Result")
        fill = FILL_HIT if result == "HIT" else FILL_MISS if result == "MISS" else None
        if fill:
            for cidx in range(1, len(HEADERS) + 1):
                ws.cell(row=row, column=cidx).fill = fill
    _rebuild_track_record(wb)
    wb.save(config.EXCEL_PATH)


def _num(v):
    return float(v) if isinstance(v, (int, float)) else None


def _rebuild_track_record(wb: Workbook) -> None:
    ws = wb["Picks"]
    rows = []
    for row in range(2, ws.max_row + 1):
        rows.append({h: ws.cell(row=row, column=COL[h]).value for h in HEADERS})
    sb = wb["Track Record"]
    sb.delete_rows(1, sb.max_row + 1)

    def avg(vals):
        vals = [v for v in vals if v is not None]
        return round(sum(vals) / len(vals), 1) if vals else ""

    def stats(sub):
        decided = [r for r in sub if r["Result"] in ("HIT", "MISS")]
        hits = [r for r in decided if r["Result"] == "HIT"]
        return {
            "picks": len(sub), "open": sum(1 for r in sub if r["Result"] == "OPEN"),
            "decided": len(decided), "hits": len(hits),
            "hit_rate": round(100 * len(hits) / len(decided), 1) if decided else "",
            "pred_p5": avg(_num(r["Chance +5%"]) for r in decided),
            "pred_gain": avg(_num(r["Usual Gain %"]) for r in sub),
            "real_gain": avg(_num(r["Gain So Far %"]) for r in sub),
        }

    def emit(title, headers, data_rows):
        sb.append([title]); sb.cell(row=sb.max_row, column=1).font = Font(bold=True)
        sb.append(headers)
        for cidx in range(1, len(headers) + 1):
            sb.cell(row=sb.max_row, column=cidx).font = Font(bold=True, color="FFFFFF")
            sb.cell(row=sb.max_row, column=cidx).fill = FILL_HEAD
        for dr in data_rows:
            sb.append(dr)
        sb.append([])

    s = stats(rows)
    emit("How we're doing overall",
         ["Picks", "Waiting", "Finished", "Hit +5%", "Hit rate %",
          "We predicted %"],
         [[s["picks"], s["open"], s["decided"], s["hits"], s["hit_rate"],
           s["pred_p5"]]])
    emit("Gains — what we predicted vs what happened",
         ["Predicted usual gain %", "Actual gain so far %"],
         [[s["pred_gain"], s["real_gain"]]])

    for title, key in (("By confidence", "Confidence"), ("By list", "List"),
                       ("By engine", "Engine")):
        groups = sorted({str(r[key]) for r in rows if r[key]})
        emit(title, [key, "Picks", "Finished", "Hit +5%", "Hit rate %",
                     "Predicted gain %", "Actual gain %"],
             [[g] + [(x := stats([r for r in rows if str(r[key]) == g]))["picks"],
                     x["decided"], x["hits"], x["hit_rate"],
                     x["pred_gain"], x["real_gain"]] for g in groups])

    runs = sorted({str(r["Date Picked"]) for r in rows if r["Date Picked"]})
    emit("By run", ["Date", "Picks", "Waiting", "Hit +5%", "Hit rate %",
                    "Predicted gain %", "Actual gain %"],
         [[d] + [(x := stats([r for r in rows if str(r["Date Picked"]) == d]))["picks"],
                 x["open"], x["hits"], x["hit_rate"],
                 x["pred_gain"], x["real_gain"]] for d in runs])
    for colletter in ("A", "B", "C", "D", "E", "F", "G"):
        sb.column_dimensions[colletter].width = 20
